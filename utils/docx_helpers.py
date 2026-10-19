from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DOCX_HEADER_BG = "021F59"
DOCX_TEXT_COLOR = "021F59"
DOCX_ODD_ROW_BG = "F2E9DF"

def add_hyperlink(paragraph, url, text=None):
    """Append a clickable hyperlink (blue, underlined) to a paragraph."""
    text = text or url
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    run = OxmlElement('w:r')
    r_pr = OxmlElement('w:rPr')

    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    r_pr.append(color)
    r_pr.append(underline)

    t = OxmlElement('w:t')
    t.text = text

    run.append(r_pr)
    run.append(t)
    #run.font.name = "Arial"

    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink

def write_cell(cell, text, font_size, color, bold=False,
                align=WD_ALIGN_PARAGRAPH.CENTER):
    """Write styled text into a cell."""
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run(str(text))
    run.font.size = Pt(font_size)
    run.font.name = "Arial"
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, hex_fill):
    """Set the background color of a cell (hex without '#')."""
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn('w:shd')):
        tc_pr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tc_pr.append(shd)


def add_docx_stat_table(document, heading_text, table_data, rows_def):
    """
    Generic Word table: years as columns, one row per statistic.
 
    table_data: list of dicts (one per year), must contain the key 'Year'
    rows_def:   list of (label, key, highlight_fill) tuples
                - label: text shown in the first column
                - key: key in table_data records
                - highlight_fill: hex color (no '#') for a highlighted bold
                  row, or None for a normal row (odd-row banding applies)
    """
    #rows_def = rows_def[::-1]
    table_data = table_data[::-1]
    # Heading
    heading = document.add_heading(heading_text, level=2)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x02, 0x1F, 0x59)
        run.font.name = "Roboto"
 
    n_years = len(table_data)
    table = document.add_table(rows=len(rows_def) + 1, cols=n_years + 1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
 
    # Header row: empty corner cell + one column per year
    corner = table.cell(0, 0)
    write_cell(corner, "", font_size=10, color="FFFFFF", bold=True)
    shade_cell(corner, DOCX_HEADER_BG)
    for j, record in enumerate(table_data):
        cell = table.cell(0, j + 1)
        write_cell(cell, record['Year'], font_size=10, color="FFFFFF", bold=True)
        shade_cell(cell, DOCX_HEADER_BG)
 
    # Data rows
    for i, (label, key, highlight_fill) in enumerate(rows_def):
        is_highlight = highlight_fill is not None
 
        label_cell = table.cell(i + 1, 0)
        write_cell(label_cell, label, font_size=11, color=DOCX_TEXT_COLOR,
                    bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
 
        row_cells = [label_cell]
        for j, record in enumerate(table_data):
            cell = table.cell(i + 1, j + 1)
            write_cell(cell, record[key], font_size=11,
                        color=DOCX_TEXT_COLOR, bold=is_highlight)
            row_cells.append(cell)
 
        if is_highlight:
            fill = highlight_fill
        elif i % 2 == 1:
            fill = DOCX_ODD_ROW_BG
        else:
            fill = None
        if fill:
            for cell in row_cells:
                shade_cell(cell, fill)
    keep_table_on_one_page(table)
    heading.paragraph_format.keep_with_next = True  
    return document


def add_page_number(paragraph):
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT

    page_num_run = paragraph.add_run()

    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"

    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')

    page_num_run._r.append(fldChar1)
    page_num_run._r.append(instrText)
    page_num_run._r.append(fldChar2)

    of_run = paragraph.add_run()
    t2 = OxmlElement('w:t')
    t2.set(qn('xml:space'), 'preserve')
    t2.text = '/'
    of_run._r.append(t2)

    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'begin')

    instrText2 = OxmlElement('w:instrText')
    instrText2.set(qn('xml:space'), 'preserve')
    instrText2.text = "NUMPAGES"

    fldChar4 = OxmlElement('w:fldChar')
    fldChar4.set(qn('w:fldCharType'), 'end')

    num_pages_run = paragraph.add_run()
    num_pages_run._r.append(fldChar3)
    num_pages_run._r.append(instrText2)
    num_pages_run._r.append(fldChar4)


def keep_table_on_one_page(table):
    """Keeps a Word table in one piece: if it doesn't fit in the space left on the
    page, the whole table moves to the next page."""
    last = len(table.rows) - 1
    for i, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn('w:cantSplit')) is None:
            tr_pr.append(OxmlElement('w:cantSplit'))
        if i < last:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True